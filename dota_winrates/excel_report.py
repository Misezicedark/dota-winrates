"""Excel workbook for the kept matches. Needs openpyxl (pip install openpyxl)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import analysis
from .reports import format_date

FONT_NAME = "Arial"
RED_FILL = PatternFill("solid", start_color="FFC7CE")
GREEN_FILL = PatternFill("solid", start_color="C6EFCE")
GRAY_FILL = PatternFill("solid", start_color="EDEDED")
HEADER_FILL = PatternFill("solid", start_color="1F3864")
TOTAL_FILL = PatternFill("solid", start_color="DDEBF7")
MY_TEAM_FILL = PatternFill("solid", start_color="2E75B6")
ENEMY_FILL = PatternFill("solid", start_color="C55A11")
THIN_BORDER = Border(bottom=Side(style="thin", color="BFBFBF"))
PERCENT_FORMAT = "0.0%"
DIFFERENCE_FORMAT = "+0.0;-0.0;0.0"
STRATZ_COLOR = "7030A0"
LINK_COLOR = "0563C1"
SOURCE_LABELS = {
    analysis.SOURCE_OPENDOTA: "OpenDota ranked",
    analysis.SOURCE_STRATZ: "Stratz total",
    analysis.SOURCE_ANONYMOUS: "anonymous",
    analysis.SOURCE_UNKNOWN: "unknown",
}
MATCHES_SHEET = "Matches"
PLAYERS_SHEET = "Players"
SUMMARY_SHEET = "Summary"
FIRST_MATCH_ROW = 4
TEAM_SIZE = 5


def font(bold=False, color="000000", size=10, italic=False, underline=None):
    return Font(name=FONT_NAME, bold=bold, color=color, size=size, italic=italic, underline=underline)


def set_text(cell, text: str) -> None:
    """Writes text that stays text even when it starts with '=' (player names are user-chosen)."""
    cell.value = text
    cell.data_type = "s"


def style_header(sheet, row: int, first_column: int, last_column: int, fill) -> None:
    for column in range(first_column, last_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_players_sheet(workbook, kept: list[analysis.MatchRecord]) -> dict[tuple[int, str, int], int]:
    """Writes one row per player of every kept match; returns (match id, team, index) → row number."""
    sheet = workbook.create_sheet(PLAYERS_SHEET)
    headers = ["Match", "Series", "Date", "Team", "Name", "Hero", "Wins", "Losses", "Winrate", "Source", "Account ID"]
    sheet.append(headers)
    style_header(sheet, 1, 1, len(headers), HEADER_FILL)
    row_numbers = {}
    row = 2
    for record in kept:
        for team_key, team_label, team in (("my_team", "My team", record.my_team), ("enemy", "Enemy", record.enemy_team)):
            for index, player in enumerate(team):
                is_stratz = player.winrate.source == analysis.SOURCE_STRATZ
                values = [
                    record.match_id, record.series, format_date(record.start_time), team_label, None, None,
                    player.winrate.wins, player.winrate.losses,
                    f'=IF(AND(ISNUMBER(G{row}),ISNUMBER(H{row})),IF(G{row}+H{row}>0,G{row}/(G{row}+H{row}),"-"),"-")',
                    SOURCE_LABELS[player.winrate.source], player.account_id,
                ]
                for column, value in enumerate(values, start=1):
                    cell = sheet.cell(row=row, column=column, value=value)
                    cell.font = font(bold=player.is_me, italic=is_stratz, color=STRATZ_COLOR if is_stratz else "000000")
                    cell.border = THIN_BORDER
                set_text(sheet.cell(row=row, column=5), player.name + (" (you)" if player.is_me else ""))
                set_text(sheet.cell(row=row, column=6), player.hero)
                sheet.cell(row=row, column=9).number_format = PERCENT_FORMAT
                sheet.cell(row=row, column=11).number_format = "0"
                row_numbers[(record.match_id, team_key, index)] = row
                row += 1
    for column, width in enumerate([13, 7, 17, 10, 24, 20, 9, 9, 10, 16, 13], start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    return row_numbers


def build_matches_sheet(workbook, kept: list[analysis.MatchRecord], row_numbers: dict) -> int:
    """Writes one coloured row per kept match; returns the last data row (at least the first data row)."""
    sheet = workbook.active
    sheet.title = MATCHES_SHEET
    sheet["A1"] = (
        "Green = my team's average is higher · Red = the enemy team's average is higher or equal · "
        "Purple italic winrate = Stratz total (hidden match history) · hover a winrate to see the player"
    )
    sheet["A1"].font = font(bold=True)
    sheet.merge_cells("A1:U1")
    sheet["G2"] = "My team"
    sheet.merge_cells("G2:M2")
    sheet["N2"] = "Enemy team"
    sheet.merge_cells("N2:T2")
    style_header(sheet, 2, 7, 13, MY_TEAM_FILL)
    style_header(sheet, 2, 14, 20, ENEMY_FILL)
    headers = (
        ["#", "Series", "Match", "Date", "My hero", "Result", "Me"]
        + [f"Teammate {number}" for number in range(1, TEAM_SIZE)]
        + ["My team avg", "Known"]
        + [f"Enemy {number}" for number in range(1, TEAM_SIZE + 1)]
        + ["Enemy avg", "Known", "Difference (pp)"]
    )
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=column, value=header)
    style_header(sheet, 3, 1, len(headers), HEADER_FILL)
    sheet.row_dimensions[3].height = 30

    for position, record in enumerate(kept):
        row = FIRST_MATCH_ROW + position
        sheet.cell(row=row, column=1, value=position + 1)
        sheet.cell(row=row, column=2, value=record.series)
        sheet.cell(row=row, column=3, value=record.match_id).hyperlink = f"https://www.opendota.com/matches/{record.match_id}"
        sheet.cell(row=row, column=4, value=format_date(record.start_time))
        set_text(sheet.cell(row=row, column=5), record.hero)
        sheet.cell(row=row, column=6, value="Win" if record.won else "Loss")

        # "Me" first, then teammates in slot order (sorted() is stable).
        my_order = sorted(range(len(record.my_team)), key=lambda index: not record.my_team[index].is_me)
        cells = [(7 + offset, "my_team", index, record.my_team[index]) for offset, index in enumerate(my_order[:TEAM_SIZE])]
        cells += [(14 + index, "enemy", index, player) for index, player in enumerate(record.enemy_team[:TEAM_SIZE])]
        stratz_columns = set()
        for column, team_key, index, player in cells:
            cell = sheet.cell(row=row, column=column, value=f"='{PLAYERS_SHEET}'!I{row_numbers[(record.match_id, team_key, index)]}")
            cell.number_format = PERCENT_FORMAT
            cell.alignment = Alignment(horizontal="right")
            record_text = "no data" if player.winrate.wins is None else f"{player.winrate.wins} W / {player.winrate.losses} L"
            comment = Comment(f"{player.name}\n{player.hero}\n{SOURCE_LABELS[player.winrate.source]}: {record_text}", "dota-winrates")
            comment.width, comment.height = 200, 75
            cell.comment = comment
            if player.winrate.source == analysis.SOURCE_STRATZ:
                stratz_columns.add(column)

        sheet.cell(row=row, column=12, value=f'=IF(COUNT(G{row}:K{row})=0,"",AVERAGE(G{row}:K{row}))').number_format = PERCENT_FORMAT
        sheet.cell(row=row, column=13, value=f"=COUNT(G{row}:K{row})")
        sheet.cell(row=row, column=19, value=f'=IF(COUNT(N{row}:R{row})=0,"",AVERAGE(N{row}:R{row}))').number_format = PERCENT_FORMAT
        sheet.cell(row=row, column=20, value=f"=COUNT(N{row}:R{row})")
        sheet.cell(row=row, column=21, value=f'=IF(AND(ISNUMBER(L{row}),ISNUMBER(S{row})),(L{row}-S{row})*100,"")').number_format = DIFFERENCE_FORMAT

        row_fill = GREEN_FILL if record.difference > 0 else RED_FILL
        for column in range(1, len(headers) + 1):
            cell = sheet.cell(row=row, column=column)
            cell.fill = row_fill
            cell.border = THIN_BORDER
            if column == 3:
                cell.font = font(color=LINK_COLOR, underline="single")
            elif column in stratz_columns:
                cell.font = font(italic=True, color=STRATZ_COLOR)
            else:
                cell.font = font(bold=column in (12, 19, 21))
        for column in (2, 13, 20):
            sheet.cell(row=row, column=column).alignment = Alignment(horizontal="center")

    widths = [5, 7, 12, 17, 17, 8, 8] + [10] * 4 + [11, 8] + [9] * 5 + [10, 8, 11]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "G4"
    last_row = max(FIRST_MATCH_ROW, FIRST_MATCH_ROW + len(kept) - 1)
    sheet.auto_filter.ref = f"A3:U{last_row}"
    return last_row


def build_summary_sheet(workbook, records: list[analysis.MatchRecord], last_row: int, min_difference: float) -> None:
    sheet = workbook.create_sheet(SUMMARY_SHEET, index=0)

    def column_range(column: str) -> str:
        return f"{MATCHES_SHEET}!${column}${FIRST_MATCH_ROW}:${column}${last_row}"

    series_range, result_range, difference_range = column_range("B"), column_range("F"), column_range("U")
    sheet["A1"] = f"Solo ranked matches: my team's average winrate vs the enemy team's (difference of at least {min_difference:g} pp)"
    sheet["A1"].font = font(bold=True, size=14)
    headers = ["Series", "Period", "Kept matches", "Wins", "Win %", "We higher", "of which won", "% won",
               "They higher", "of which won", "% won"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=column, value=header)
    style_header(sheet, 3, 1, len(headers), HEADER_FILL)
    sheet.row_dimensions[3].height = 30

    summaries = analysis.summarize(records)
    for offset, summary in enumerate(summaries):
        row = 4 + offset
        prefix = "" if summary.series is None else f"{series_range},{summary.series},"

        def count(criteria: str) -> str:
            return f"=COUNTIFS({prefix}{criteria})"

        period = ""
        if summary.first_start_time is not None:
            period = f"{format_date(summary.first_start_time)[:10]} → {format_date(summary.last_start_time)[:10]}"
        # Match ids are numbers, so COUNT gives the kept matches of the total row.
        kept_formula = f"=COUNT({column_range('C')})" if summary.series is None else f"=COUNTIFS({series_range},{summary.series})"
        values = [
            "Total" if summary.series is None else f"Series {summary.series}",
            period,
            kept_formula,
            count(f'{result_range},"Win"'),
            f'=IF(C{row}=0,"",D{row}/C{row})',
            count(f'{difference_range},">0"'),
            count(f'{difference_range},">0",{result_range},"Win"'),
            f'=IF(F{row}=0,"",G{row}/F{row})',
            count(f'{difference_range},"<=0"'),
            count(f'{difference_range},"<=0",{result_range},"Win"'),
            f'=IF(I{row}=0,"",J{row}/I{row})',
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = font(bold=summary.series is None)
            cell.border = THIN_BORDER
            if summary.series is None:
                cell.fill = TOTAL_FILL
        for column in (5, 8, 11):
            sheet.cell(row=row, column=column).number_format = "0%"

    total_row = 4 + len(summaries) - 1
    overall_row = total_row + 2
    sheet.cell(row=overall_row, column=1, value="Overall").font = font(bold=True, size=12)
    overall_lines = [
        ("We had the higher average", f"=F{total_row}", GREEN_FILL),
        ("They had the higher average", f"=I{total_row}", RED_FILL),
    ]
    for offset, (label, count_formula, fill) in enumerate(overall_lines, start=1):
        row = overall_row + offset
        values = [label, count_formula, f'=IF(C{total_row}=0,"",B{row}/C{total_row})']
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = font(bold=True)
            cell.fill = fill
        sheet.cell(row=row, column=3).number_format = PERCENT_FORMAT

    notes = [
        "Notes",
        "Only solo matchmaking matches (my party size = 1) where both teams have at least one known winrate"
        f" and the team averages differ by at least {min_difference:g} percentage points.",
        "Public players: ranked wins / (wins + losses), whole career (OpenDota).",
        "Hidden match history (purple, italic): wins / matches across all game modes (Stratz).",
        "Anonymous players have no account id in the match data and are left out of the averages.",
        "Team average includes you. Difference = my team average - enemy team average.",
        "Matches by step: " + ", ".join(f"{label} {count}" for label, count in analysis.exclusion_funnel(records)) + ".",
    ]
    notes_row = overall_row + len(overall_lines) + 2
    for offset, note in enumerate(notes):
        sheet.cell(row=notes_row + offset, column=1, value=note).font = font(bold=offset == 0)
    for column, width in zip("ABCDEFGHIJK", [30, 26, 12, 8, 8, 10, 11, 8, 11, 11, 8]):
        sheet.column_dimensions[column].width = width


def write_workbook(records: list[analysis.MatchRecord], path: Path, min_difference: float) -> None:
    kept = [record for record in records if record.included]
    workbook = Workbook()
    row_numbers = build_players_sheet(workbook, kept)
    last_row = build_matches_sheet(workbook, kept, row_numbers)
    build_summary_sheet(workbook, records, last_row, min_difference)
    workbook.active = 0
    # openpyxl stores formulas without results; this makes Excel and Numbers compute them on open.
    workbook.calculation.fullCalcOnLoad = True
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
