import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

try:
    import openpyxl
    from pycel import ExcelCompiler
except ImportError:
    openpyxl = None

from dota_winrates.analysis import (
    EXCLUDED_TOO_CLOSE,
    PARTY_SOURCE_OPENDOTA,
    SOURCE_OPENDOTA,
    SOURCE_STRATZ,
    MatchRecord,
    PlayerRow,
    PlayerWinrate,
)

GREEN = "C6EFCE"
RED = "FFC7CE"


def team(winrates, include_me=False, stratz_index=None, first_account=1):
    rows = []
    for index, winrate in enumerate(winrates):
        is_me = include_me and index == 0
        source = SOURCE_STRATZ if index == stratz_index else SOURCE_OPENDOTA
        wins = round(winrate * 100)
        rows.append(PlayerRow(103093152 if is_me else first_account + index, "Dark" if is_me else f"Player {first_account + index}",
                              "Lion", is_me, PlayerWinrate(winrate, source, wins, 100 - wins)))
    return tuple(rows)


def match(match_id, series, my_winrates, enemy_winrates, won, exclusion_reason=None, stratz_index=None):
    my_team = team(my_winrates, include_me=True, stratz_index=stratz_index, first_account=10)
    enemy_team = team(enemy_winrates, first_account=20)
    average = lambda rows: sum(row.winrate.winrate for row in rows) / len(rows)
    return MatchRecord(match_id=match_id, series=series, start_time=1_790_000_000 - match_id, hero="Lion", won=won,
                       party_size=1, party_source=PARTY_SOURCE_OPENDOTA, exclusion_reason=exclusion_reason,
                       my_team=my_team, enemy_team=enemy_team, my_average=average(my_team), enemy_average=average(enemy_team))


@unittest.skipIf(openpyxl is None, "needs openpyxl and pycel (see the development setup in README)")
class ExcelReportTest(unittest.TestCase):
    def setUp(self):
        from dota_winrates.excel_report import write_workbook

        self.temporary_directory = tempfile.TemporaryDirectory()
        self.path = Path(self.temporary_directory.name) / "report.xlsx"
        self.records = [
            match(1, 1, [0.5, 0.5, 0.5, 0.5, 0.5], [0.6, 0.6, 0.6, 0.6, 0.6], won=True),           # they higher, won
            match(2, 1, [0.6, 0.6, 0.6, 0.6, 0.6], [0.5, 0.5, 0.5, 0.5, 0.5], won=False, stratz_index=1),  # we higher, lost
            match(3, 2, [0.5, 0.5, 0.5, 0.5, 0.5], [0.55, 0.55, 0.55, 0.55, 0.55], won=False),     # they higher, lost
            match(4, 2, [0.5] * 5, [0.501] * 5, won=True, exclusion_reason=EXCLUDED_TOO_CLOSE),   # excluded
        ]
        write_workbook(self.records, self.path, 1.0)
        self.workbook = openpyxl.load_workbook(self.path)
        self.compiler = ExcelCompiler(filename=str(self.path))

    def tearDown(self):
        self.temporary_directory.cleanup()

    def value(self, reference):
        return self.compiler.evaluate(reference)

    def test_has_the_three_sheets_with_summary_first(self):
        self.assertEqual(self.workbook.sheetnames, ["Summary", "Matches", "Players"])

    def test_matches_sheet_lists_only_kept_matches(self):
        sheet = self.workbook["Matches"]
        self.assertEqual([sheet.cell(row=row, column=3).value for row in range(4, sheet.max_row + 1)], [1, 2, 3])

    def test_rows_are_coloured_by_which_team_was_higher(self):
        sheet = self.workbook["Matches"]
        fills = [sheet.cell(row=row, column=1).fill.start_color.rgb[-6:] for row in (4, 5, 6)]
        self.assertEqual(fills, [RED, GREEN, RED])

    def test_difference_formula_gives_points(self):
        self.assertAlmostEqual(self.value("Matches!U4"), -10.0)
        self.assertAlmostEqual(self.value("Matches!U5"), 10.0)

    def test_my_column_comes_first_and_resolves_my_winrate(self):
        self.assertAlmostEqual(self.value("Matches!G4"), 0.5)

    def test_summary_counts_per_series_and_total(self):
        series_one = [self.value(f"Summary!{column}4") for column in "CDFGIJ"]
        total = [self.value(f"Summary!{column}6") for column in "CDFGIJ"]
        self.assertEqual(series_one, [2, 1, 1, 0, 1, 1])
        self.assertEqual(total, [3, 1, 1, 0, 2, 1])

    def test_overall_block_shows_shares_of_kept_matches(self):
        # Summary rows: 4 = series 1, 5 = series 2, 6 = total, 8 = "Overall", 9 = we higher, 10 = they higher.
        self.assertAlmostEqual(self.value("Summary!C9"), 1 / 3)
        self.assertAlmostEqual(self.value("Summary!C10"), 2 / 3)

    def test_stratz_winrates_are_purple_italic(self):
        cell = self.workbook["Matches"].cell(row=5, column=8)
        self.assertTrue(cell.font.italic)
        self.assertEqual(cell.font.color.rgb[-6:], "7030A0")

    def test_formula_like_player_names_stay_text(self):
        from dota_winrates.excel_report import write_workbook

        record = self.records[0]
        renamed = replace(record.my_team[1], name="=1+1")
        record = replace(record, my_team=(record.my_team[0], renamed, *record.my_team[2:]))
        write_workbook([record], self.path, 1.0)
        players = openpyxl.load_workbook(self.path)["Players"]
        name_cell = players.cell(row=3, column=5)
        self.assertEqual((name_cell.value, name_cell.data_type), ("=1+1", "s"))

    def test_workbook_without_kept_matches_is_still_valid(self):
        from dota_winrates.excel_report import write_workbook

        write_workbook([self.records[3]], self.path, 1.0)
        self.assertEqual(ExcelCompiler(filename=str(self.path)).evaluate("Summary!C4"), 0)


if __name__ == "__main__":
    unittest.main()
